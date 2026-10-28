from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from .models import Message
from .services import add_member, create_channel, get_or_create_dm

pytestmark = pytest.mark.django_db


def sheets(response):
    book = load_workbook(BytesIO(response.content))
    return {name: list(book[name].values) for name in book.sheetnames}


def say(conversation, author, body):
    return Message.objects.create(conversation=conversation, author=author, body=body)


def test_member_exports_the_whole_conversation_not_just_the_latest_page(client, alice):
    channel = create_channel(alice, "ops", is_private=False)
    for n in range(205):
        say(channel, alice, f"message {n}")
    say(channel, alice, "=SUM(1)")
    client.force_login(alice)
    response = client.get(reverse("chat:export_conversation", args=[channel.pk]))
    data = sheets(response)
    assert len(data["Messages"]) == 1 + 1 + 205 + 1  # header, "joined", the rest
    assert data["Messages"][-1][4] == "'=SUM(1)"
    assert data["Conversations"][1][0] == "#ops"
    assert "ops" in response["Content-Disposition"]


def test_dm_title_is_the_other_person_and_system_lines_name_the_subject(
    client, alice, bob
):
    dm = get_or_create_dm(alice, bob)
    say(dm, bob, "hi")
    client.force_login(alice)
    data = sheets(client.get(reverse("chat:export_conversation", args=[dm.pk])))
    assert data["Messages"][1][0] == "bob"
    private = create_channel(alice, "boardroom", is_private=True)
    add_member(private, alice, bob)
    data = sheets(client.get(reverse("chat:export_conversation", args=[private.pk])))
    assert data["Messages"][-1][5] == "bob"


def test_a_non_member_cannot_export_a_public_channel_or_a_private_one(
    client, alice, bob
):
    public = create_channel(alice, "open", is_private=False)
    private = create_channel(alice, "secret", is_private=True)
    client.force_login(bob)
    assert (
        client.get(reverse("chat:export_conversation", args=[public.pk])).status_code
        == 403
    )
    assert (
        client.get(reverse("chat:export_conversation", args=[private.pk])).status_code
        == 404
    )


def test_export_all_covers_only_my_conversations(client, alice, bob, carol):
    mine = create_channel(alice, "mine", is_private=True)
    say(mine, alice, "for me")
    theirs = get_or_create_dm(bob, carol)
    say(theirs, bob, "not for alice")
    client.force_login(alice)
    data = sheets(client.get(reverse("chat:export")))
    bodies = [row[4] for row in data["Messages"][1:]]
    assert "for me" in bodies
    assert "not for alice" not in bodies
