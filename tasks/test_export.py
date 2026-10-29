from datetime import timedelta
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from .services import add_update, reschedule

pytestmark = pytest.mark.django_db


def sheets(response):
    book = load_workbook(BytesIO(response.content))
    return {name: list(book[name].values) for name in book.sheetnames}


def test_admin_exports_every_task_with_updates(client, boss, alice, task):
    add_update(task, alice, "=HYPERLINK(1)")
    client.force_login(boss)
    data = sheets(client.get(reverse("tasks:export")))
    assert data["Tasks"][1][:2] == (task.pk, "alice")
    assert data["Updates"][1][3] == "'=HYPERLINK(1)"


def test_date_changes_are_exported_with_their_reason(client, boss, task):
    reschedule(task, boss, task.due_date + timedelta(days=2), "=Waiting")
    client.force_login(boss)
    data = sheets(client.get(reverse("tasks:export")))
    row = data["Date changes"][1]
    assert row[0] == task.pk
    assert row[5] == "'=Waiting"


def test_employee_exports_only_their_own(client, bob, task):
    client.force_login(bob)
    data = sheets(client.get(reverse("tasks:export")))
    assert len(data["Tasks"]) == 1


def test_anonymous_is_sent_to_login(client):
    assert client.get(reverse("tasks:export")).status_code == 302
