from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from core.models import AccessGrant, Company
from people.models import Person

from .models import Requisition, SourcedNumber
from .services import source_number

pytestmark = pytest.mark.django_db


@pytest.fixture
def senken():
    return Company.objects.get(slug="senken")


@pytest.fixture
def caller(client, django_user_model, senken):
    user = django_user_model.objects.create_user(username="S9", password="pw")
    senken.memberships.create(user=user)
    for name in ("candidates", "openings"):
        AccessGrant.objects.create(
            user=user, company=senken, module="recruitment", section=name
        )
    client.force_login(user)
    return user


@pytest.fixture
def req(senken, caller):
    return Requisition.objects.create(
        company=senken, opened_by=caller, department="bo", profile="Fitter", headcount=3
    )


def sheet(*rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = BytesIO()
    book.save(out)
    return SimpleUploadedFile("n.xlsx", out.getvalue())


def test_new_number_makes_a_card_and_repeat_links_to_it(req, caller):
    a = source_number(req, "Asha", "+91 98765 43210", caller, "Naukri")
    b = source_number(req, "Asha K", "09876543210", caller, "Referral")
    assert (a.is_duplicate, b.is_duplicate) == (False, True)
    assert b.person == a.person and Person.objects.count() == 1
    assert SourcedNumber.objects.count() == 2


def test_short_number_rejected(req, caller):
    with pytest.raises(ValueError):
        source_number(req, "X", "12345", caller)
    assert not SourcedNumber.objects.exists()


def test_add_form_and_counts_on_requisition_list(client, req):
    for phone in ("9876543210", "9876543210", "9123456789"):
        r = client.post(
            reverse("hiring:source_new"),
            {"requisition": req.pk, "name": "A", "phone": phone, "source": "Naukri"},
        )
        assert r.status_code == 302
    assert Person.objects.count() == 2
    page = client.get(reverse("hiring:list")).content.decode()
    assert "3 sourced of 3 needed" in page and "1 duplicate" in page


def test_closed_requisition_cannot_be_sourced(client, req):
    req.status = "closed"
    req.save()
    r = client.post(
        reverse("hiring:source_new"),
        {"requisition": req.pk, "name": "A", "phone": "9876543210"},
    )
    assert r.status_code == 200 and not SourcedNumber.objects.exists()


def test_import_counts_new_duplicates_and_skips(client, req):
    upload = sheet(
        ("Name", "Mobile", "Source"),
        ("Asha", 9876543210.0, "Naukri"),
        ("Asha again", "+91 98765 43210", None),
        ("Bad", "123", None),
        (None, "9123456789", None),
        ("Ravi", "9123456780", None),
    )
    r = client.post(
        reverse("hiring:source_import"),
        {"requisition": req.pk, "source": "Walk-in", "sheet": upload},
        follow=True,
    )
    assert "2 new, 1 duplicate. Skipped 2" in r.content.decode()
    assert list(
        SourcedNumber.objects.order_by("id").values_list("source", flat=True)
    ) == ["Naukri", "Walk-in", "Walk-in"]


def test_import_without_headings_is_refused(client, req):
    r = client.post(
        reverse("hiring:source_import"),
        {"requisition": req.pk, "sheet": sheet(("Asha", "9876543210"))},
    )
    assert b"Name and Phone headings" in r.content
    assert not SourcedNumber.objects.exists()


def test_filter_and_export(client, req, caller):
    source_number(req, "A", "9876543210", caller, "Naukri")
    source_number(req, "A", "9876543210", caller, "Naukri")
    page = client.get(reverse("hiring:sourcing") + "?dup=1").content.decode()
    assert page.count(">Duplicate<") == 1
    r = client.get(reverse("hiring:sourcing_export"))
    book = load_workbook(BytesIO(r.content))
    assert list(book["By source"].values)[1] == ("Naukri", 2, 1, 0.5)


def test_other_company_is_hidden(client, req, caller):
    other = Company.objects.exclude(slug="senken").first()
    foreign = Requisition.objects.create(
        company=other, opened_by=caller, department="bo", profile="Zed", headcount=1
    )
    source_number(foreign, "Hidden", "9876543210", caller)
    assert b"Hidden" not in client.get(reverse("hiring:sourcing")).content
