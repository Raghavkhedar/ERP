from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.db import IntegrityError
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import AccessGrant, Company

from .models import FollowUp, Offer, Requisition
from .services import (
    OfferError,
    accept,
    add_round,
    decline,
    log_call,
    record_interview,
    source_number,
    start_offer,
)

pytestmark = pytest.mark.django_db


@pytest.fixture
def senken():
    return Company.objects.get(slug="senken")


@pytest.fixture
def caller(client, django_user_model, senken):
    user = django_user_model.objects.create_user(username="S9", password="pw")
    senken.memberships.create(user=user)
    AccessGrant.objects.create(
        user=user, company=senken, module="recruitment", section="candidates"
    )
    client.force_login(user)
    return user


@pytest.fixture
def waiting(senken, caller):
    req = Requisition.objects.create(
        company=senken, opened_by=caller, department="bo", profile="Fitter", headcount=2
    )
    source_number(req, "Asha", "9876543210", caller)
    log_call(
        FollowUp.objects.get(),
        "interview",
        caller,
        next_due=timezone.now() + timedelta(days=1),
    )
    record_interview(FollowUp.objects.get(status="pending"), "selected", caller)
    return FollowUp.objects.get(status="pending", purpose="offer")


@pytest.fixture
def offer(waiting, caller):
    return start_offer(waiting, Decimal("30000"), caller)


def test_start_offer_completes_the_follow_up_and_logs_round_one(waiting, caller):
    offer = start_offer(waiting, Decimal("30000"), caller)
    waiting.refresh_from_db()
    assert waiting.status == "completed"
    assert (offer.state, offer.latest.amount) == ("offered", Decimal("30000"))


def test_rounds_track_the_negotiation(offer, caller):
    add_round(
        offer, Decimal("32000"), caller, counter=Decimal("35000"), note="asked more"
    )
    assert offer.rounds.count() == 2
    assert offer.latest.counter == Decimal("35000")


def test_accept_books_the_joining(offer, caller):
    accept(offer, date(2030, 1, 6), caller)
    offer.refresh_from_db()
    assert (offer.state, offer.joining_date) == ("accepted", date(2030, 1, 6))
    joining = FollowUp.objects.get(purpose="joining")
    assert joining.due.date() == date(2030, 1, 6)


def test_a_closed_offer_cannot_change_again(offer, caller):
    decline(offer, "better offer elsewhere")
    for act in (
        lambda: add_round(offer, Decimal("1"), caller),
        lambda: accept(offer, date(2030, 1, 6), caller),
        lambda: decline(offer),
    ):
        with pytest.raises(OfferError):
            act()
    assert FollowUp.objects.filter(purpose="joining").count() == 0


def test_only_one_open_offer_per_person(offer, caller):
    with pytest.raises(IntegrityError):
        Offer.objects.create(person=offer.person, made_by=caller)


def test_pages_prefill_and_flow(client, waiting, caller):
    waiting.person.expected_salary = Decimal("28000")
    waiting.person.save()
    url = reverse("hiring:offer_start", args=[waiting.pk])
    assert b'value="28000' in client.get(url).content
    r = client.post(url, {"amount": "28000"})
    offer = Offer.objects.get()
    assert r.status_code == 302 and r.url == reverse("hiring:offer", args=[offer.pk])
    client.post(
        reverse("hiring:offer_round", args=[offer.pk]),
        {"amount": "30000", "counter": "33000"},
    )
    client.post(
        reverse("hiring:offer_accept", args=[offer.pk]), {"joining_date": "2030-01-06"}
    )
    offer.refresh_from_db()
    assert offer.state == "accepted" and offer.rounds.count() == 2
    assert b"Accepted" in client.get(reverse("hiring:offers")).content


def test_export(client, offer):
    book = load_workbook(BytesIO(client.get(reverse("hiring:offers_export")).content))
    assert list(book["Offers"].values)[1][3:5] == ("Offered", 30000.0)
    assert len(list(book["Rounds"].values)) == 2


def test_marking_joined_closes_the_joining_follow_up(client, offer, caller):
    accept(offer, date(2030, 1, 6), caller)
    r = client.post(
        reverse("hiring:joined", args=[offer.pk]), {"joined_on": "2030-01-07"}
    )
    assert r.status_code == 302
    offer.refresh_from_db()
    assert offer.joined_on == date(2030, 1, 7)
    assert not FollowUp.objects.filter(purpose="joining", status="pending").exists()
    page = client.get(reverse("hiring:onboarding")).content.decode()
    assert "joined 7 Jan 2030" in page and "coming soon" in page


def test_only_accepted_unjoined_offers_can_join(client, offer, caller):
    client.post(reverse("hiring:joined", args=[offer.pk]), {"joined_on": "2030-01-07"})
    offer.refresh_from_db()
    assert offer.joined_on is None
    accept(offer, date(2030, 1, 6), caller)
    for _ in range(2):
        client.post(
            reverse("hiring:joined", args=[offer.pk]), {"joined_on": "2030-01-07"}
        )
    offer.refresh_from_db()
    assert offer.joined_on == date(2030, 1, 7)


def test_onboarding_export(client, offer, caller):
    accept(offer, date(2030, 1, 6), caller)
    book = load_workbook(
        BytesIO(client.get(reverse("hiring:onboarding_export")).content)
    )
    assert list(book["On-Boarding"].values)[1][0] == "Asha"
