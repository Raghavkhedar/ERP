from django.db import transaction
from openpyxl import load_workbook

from .services import source_number

MAX_BYTES = 1_000_000
MAX_ROWS = 2000

HEADS = {
    "name": {"name", "candidate", "candidate name", "full name"},
    "phone": {"phone", "mobile", "mobile no", "phone no", "contact", "number"},
    "email": {"email", "e-mail", "email id"},
    "source": {"source", "keyword", "source keyword"},
}


class SheetError(Exception):
    pass


def _text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 9876543210 as 9876543210.0
    return "" if value is None else str(value).strip()


def import_numbers(requisition, upload, user, source=""):
    """Source every row of an .xlsx with name and phone columns (email and source
    optional). Returns (new, duplicates, skipped) where skipped is a list of
    (row number, reason). Rows are all-or-nothing only for unreadable files."""
    if upload.size > MAX_BYTES:
        raise SheetError("That file is too big. Keep it under 1 MB.")
    try:
        sheet = load_workbook(upload, read_only=True, data_only=True).active
        rows = [
            list(r) for r in sheet.iter_rows(max_row=MAX_ROWS + 1, values_only=True)
        ]
    except Exception:
        raise SheetError(
            "Couldn't read that file. Upload an Excel .xlsx file."
        ) from None

    heads = [_text(c).lower() for c in rows[0]] if rows else []
    cols = {
        k: next((i for i, h in enumerate(heads) if h in v), None)
        for k, v in HEADS.items()
    }
    if cols["name"] is None or cols["phone"] is None:
        raise SheetError("The first row needs Name and Phone headings.")

    def cell(row, key):
        i = cols[key]
        return _text(row[i]) if i is not None and i < len(row) else ""

    new = duplicates = 0
    skipped = []
    with transaction.atomic():
        for n, row in enumerate(rows[1:], start=2):
            if not any(c not in (None, "") for c in row):
                continue
            if not cell(row, "name"):
                skipped.append((n, "no name"))
                continue
            try:
                made = source_number(
                    requisition,
                    cell(row, "name"),
                    cell(row, "phone"),
                    user,
                    source=cell(row, "source") or source,
                    email=cell(row, "email"),
                )
            except ValueError as exc:
                skipped.append((n, str(exc)))
                continue
            duplicates += made.is_duplicate
            new += not made.is_duplicate
    return new, duplicates, skipped
