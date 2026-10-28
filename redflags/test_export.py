from datetime import date
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from .services import add_note, assign_to_supervisor, submit

pytestmark = pytest.mark.django_db


def sheets(response):
    book = load_workbook(BytesIO(response.content))
    return {name: list(book[name].values) for name in book.sheetnames}


@pytest.fixture
def lead(django_user_model):
    return django_user_model.objects.create_user(
        username="lead", first_name="lead", password="pw"
    )


@pytest.fixture
def flag(alice, bob, lead):
    bob.supervisor = lead
    bob.save()
    return submit(alice, bob, date(2026, 10, 1), "=Shouted at a customer.")


def test_superadmin_exports_flags_with_notes(client, boss, flag):
    add_note(flag, boss, "Spoke to them")
    client.force_login(boss)
    data = sheets(client.get(reverse("redflags:export") + "?show=all"))
    row = data["Red flags"][1]
    assert row[2:4] == ("alice", "bob")
    assert row[5] == "'=Shouted at a customer."
    assert data["Notes"][1][3] == "Spoke to them"


def test_supervisor_exports_only_their_assigned_flag(client, boss, flag, lead, alice):
    submit(boss, alice, date(2026, 10, 2), "Another one.")
    assign_to_supervisor(flag, boss, "against")
    client.force_login(lead)
    data = sheets(client.get(reverse("redflags:export") + "?show=all"))
    assert [row[0] for row in data["Red flags"][1:]] == [flag.pk]


@pytest.mark.parametrize("who", ["alice", "manager", "trainer"])
def test_everyone_else_gets_a_404(client, request, who, flag):
    client.force_login(request.getfixturevalue(who))
    assert client.get(reverse("redflags:export")).status_code == 404
