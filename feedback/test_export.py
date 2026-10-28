from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from .services import add_comment, submit

pytestmark = pytest.mark.django_db


def sheets(response):
    book = load_workbook(BytesIO(response.content))
    return {name: list(book[name].values) for name in book.sheetnames}


@pytest.fixture
def item(alice):
    return submit(alice, "problem", "tasks", "=Cannot save")


def test_reader_exports_everything_with_comments_and_lessons(
    client, manager, item, alice
):
    add_comment(item, manager, "Looking into it")
    client.force_login(manager)
    data = sheets(client.get(reverse("feedback:export") + "?show=all"))
    assert data["Feedback"][1][2] == "alice"
    assert data["Feedback"][1][5] == "'=Cannot save"
    assert data["Comments"][1][3] == "Looking into it"
    assert "Lessons" in data


def test_submitter_exports_only_their_own_and_no_lessons(client, alice, bob, item):
    submit(bob, "idea", "chat", "Bob's idea")
    client.force_login(alice)
    data = sheets(client.get(reverse("feedback:export") + "?show=all"))
    assert [row[2] for row in data["Feedback"][1:]] == ["alice"]
    assert "Lessons" not in data


def test_export_follows_the_open_filter(client, manager, item):
    client.force_login(manager)
    resolved = sheets(client.get(reverse("feedback:export") + "?show=resolved"))
    assert len(resolved["Feedback"]) == 1
