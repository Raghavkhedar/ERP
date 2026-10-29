from datetime import datetime, time, timedelta
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .services import block, create_conference, respond
from .testing import book, open_slots

pytestmark = pytest.mark.django_db

DAY = timezone.localdate() + timedelta(days=1)


def at(hour):
    return timezone.make_aware(datetime.combine(DAY, time(hour)))


def sheets(response):
    book_ = load_workbook(BytesIO(response.content))
    return {name: list(book_[name].values) for name in book_.sheetnames}


def test_export_has_booked_meetings_blocked_time_and_no_open_slots(client, alice, bob):
    free, taken, held, held_too = open_slots(alice, DAY, time(9), time(13), 60)
    book(taken, bob, "=Budget talk")
    block(alice, [held.pk, held_too.pk], "Dentist")
    client.force_login(alice)
    rows = sheets(client.get(reverse("meetings:export")))["Time slots"][1:]
    assert [(r[2], r[3], r[4], r[5]) for r in rows] == [
        ("Time", "alice", "bob", "'=Budget talk"),
        ("Blocked time", "alice", None, "Dentist"),
    ]
    assert rows[1][1] - rows[1][0] == timedelta(hours=2)  # the two slots are one row


def test_guest_sees_their_booking_but_a_stranger_sees_none(client, alice, bob, carol):
    [slot] = open_slots(alice, DAY, time(9), time(10), 60)
    book(slot, bob, "Review")
    client.force_login(bob)
    assert len(sheets(client.get(reverse("meetings:export")))["Time slots"]) == 2
    client.force_login(carol)
    assert len(sheets(client.get(reverse("meetings:export")))["Time slots"]) == 1


def test_conferences_list_attendees_and_my_answer(client, alice, bob, carol):
    conference = create_conference(
        alice, "Quarterly review", at(10), at(11), [bob, carol], location="Room 2"
    )
    respond(conference, bob, True)
    client.force_login(bob)
    data = sheets(client.get(reverse("meetings:export")))
    assert data["Conferences"][1][1] == "Quarterly review"
    assert data["Conferences"][1][8] == "Accepted"
    answers = {row[3]: row[4] for row in data["Attendees"][1:]}
    assert answers == {"bob": "Accepted", "carol": "Invited"}
    client.force_login(alice)
    assert sheets(client.get(reverse("meetings:export")))["Conferences"][1][8] == "Host"


def test_date_range_and_junk_dates(client, alice, bob):
    [slot] = open_slots(alice, DAY, time(9), time(10), 60)
    book(slot, bob, "Review")
    client.force_login(alice)
    url = reverse("meetings:export")
    after = DAY + timedelta(days=1)
    assert len(sheets(client.get(f"{url}?from={after}"))["Time slots"]) == 1
    assert len(sheets(client.get(f"{url}?from={DAY}&to={DAY}"))["Time slots"]) == 2
    assert len(sheets(client.get(f"{url}?from=nonsense"))["Time slots"]) == 2
